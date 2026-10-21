"""Bounded visual-page selection and Bedrock diagram interpretation."""

from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Protocol

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from .config import (
    BEDROCK_API_KEY,
    BEDROCK_API_MODE,
    BEDROCK_MAX_OUTPUT_TOKENS,
    BEDROCK_MODEL_ID,
    BEDROCK_MONTHLY_TOKEN_UNIT_LIMIT,
    BEDROCK_REGION,
)
from .image_utils import load_rgb_image
from .llm_provider import BedrockConfig, BedrockProvider, ProviderError
from .resource_store import ResourceStore


MAX_VISUAL_EDGE = 1_600
MAX_VISUAL_PIXELS = 2_560_000


class VisualUnderstandingError(RuntimeError):
    """A selected visual could not be interpreted safely."""


class VisualAnalyzer(Protocol):
    def describe_image(
        self, image_bytes: bytes, mime_type: str, nearby_text: str = ""
    ) -> str: ...


@dataclass(frozen=True)
class VisualCandidate:
    reference_type: str
    reference_number: int
    image_bytes: bytes
    mime_type: str = "image/jpeg"


def build_visual_analyzer(store: ResourceStore) -> BedrockProvider:
    return BedrockProvider(
        BedrockConfig(
            region=BEDROCK_REGION,
            model_id=BEDROCK_MODEL_ID,
            api_mode=BEDROCK_API_MODE,
            api_key=BEDROCK_API_KEY,
            max_tokens=BEDROCK_MAX_OUTPUT_TOKENS,
            monthly_token_unit_limit=BEDROCK_MONTHLY_TOKEN_UNIT_LIMIT,
        ),
        store,
    )


def _bounded_jpeg(image: Image.Image) -> bytes:
    image = image.convert("RGB")
    if image.width * image.height > MAX_VISUAL_PIXELS or max(image.size) > MAX_VISUAL_EDGE:
        scale = min(
            MAX_VISUAL_EDGE / max(image.size),
            (MAX_VISUAL_PIXELS / (image.width * image.height)) ** 0.5,
        )
        image = image.resize(
            (max(1, round(image.width * scale)), max(1, round(image.height * scale))),
            Image.Resampling.LANCZOS,
        )
    stream = BytesIO()
    image.save(stream, format="JPEG", quality=84, optimize=True)
    return stream.getvalue()


def _sample_numbers(numbers: list[int], maximum: int) -> list[int]:
    if len(numbers) <= maximum:
        return numbers
    if maximum == 1:
        return [numbers[len(numbers) // 2]]
    return [numbers[round(index * (len(numbers) - 1) / (maximum - 1))] for index in range(maximum)]


def _pdf_candidates(path: Path, maximum: int) -> Iterator[VisualCandidate]:
    with pymupdf.open(path) as document:
        visual_pages = [
            number
            for number, page in enumerate(document, start=1)
            if page.get_images(full=True) or page.get_drawings()
        ]
        for page_number in _sample_numbers(visual_pages, maximum):
            page = document[page_number - 1]
            zoom = min(2.0, MAX_VISUAL_EDGE / max(page.rect.width, page.rect.height))
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(zoom, zoom),
                colorspace=pymupdf.csRGB,
                alpha=False,
            )
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            yield VisualCandidate("page", page_number, _bounded_jpeg(image))


def _font() -> ImageFont.ImageFont:
    try:
        return ImageFont.truetype("arial.ttf", 17)
    except OSError:
        return ImageFont.load_default()


def _render_slide(presentation: Presentation, slide_number: int) -> bytes:
    slide = presentation.slides[slide_number - 1]
    width, height = 1_600, 900
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    font = _font()
    scale_x = width / max(1, presentation.slide_width)
    scale_y = height / max(1, presentation.slide_height)

    for shape in slide.shapes:
        left = max(0, round(shape.left * scale_x))
        top = max(0, round(shape.top * scale_y))
        right = min(width, max(left + 1, round((shape.left + shape.width) * scale_x)))
        bottom = min(height, max(top + 1, round((shape.top + shape.height) * scale_y)))
        image = getattr(shape, "image", None)
        if image is not None:
            try:
                picture = load_rgb_image(image.blob)
                picture.thumbnail((right - left, bottom - top), Image.Resampling.LANCZOS)
                canvas.paste(picture, (left, top))
            except ValueError:
                draw.rectangle((left, top, right, bottom), outline="#77775f", width=2)
            continue
        name = type(shape).__name__.casefold()
        if "connector" in name:
            draw.line((left, top, right, bottom), fill="#4f5236", width=4)
            draw.ellipse((right - 5, bottom - 5, right + 5, bottom + 5), fill="#4f5236")
            continue
        draw.rounded_rectangle(
            (left, top, right, bottom), radius=7, outline="#77775f", width=2
        )
        text = " ".join(str(getattr(shape, "text", "")).split())[:500]
        if text:
            draw.multiline_text(
                (left + 7, top + 6), text, fill="#28291d", font=font, spacing=3
            )
    return _bounded_jpeg(canvas)


def _pptx_candidates(path: Path, maximum: int) -> Iterator[VisualCandidate]:
    presentation = Presentation(path)
    visual_slides: list[int] = []
    for number, slide in enumerate(presentation.slides, start=1):
        visual_weight = 0
        for shape in slide.shapes:
            name = type(shape).__name__.casefold()
            visual_weight += int(getattr(shape, "image", None) is not None) * 3
            visual_weight += int("connector" in name) * 2
            visual_weight += int("chart" in name or "group" in name) * 3
        if visual_weight >= 2 or len(slide.shapes) >= 5:
            visual_slides.append(number)
    for slide_number in _sample_numbers(visual_slides, maximum):
        yield VisualCandidate(
            "slide", slide_number, _render_slide(presentation, slide_number)
        )


def iter_visual_candidates(
    path: Path, extension: str, maximum: int
) -> Iterator[VisualCandidate]:
    if extension == ".pdf":
        yield from _pdf_candidates(path, maximum)
    elif extension == ".pptx":
        yield from _pptx_candidates(path, maximum)
    elif extension in {".png", ".jpg", ".jpeg"}:
        yield VisualCandidate("image", 1, _bounded_jpeg(load_rgb_image(path)))


def enrich_with_visual_understanding(
    path: Path,
    extension: str,
    content: list[dict[str, object]],
    analyzer: VisualAnalyzer,
    maximum: int,
) -> list[dict[str, object]]:
    units = {
        (str(item.get("reference_type")), int(item.get("reference_number", 0))): item
        for item in content
    }
    try:
        for candidate in iter_visual_candidates(path, extension, maximum):
            unit = units.get((candidate.reference_type, candidate.reference_number))
            if unit is None:
                continue
            nearby_text = "\n".join(
                str(unit.get(field, "")) for field in ("native_text", "ocr_text")
            ).strip()
            description = analyzer.describe_image(
                candidate.image_bytes, candidate.mime_type, nearby_text
            ).strip()
            if not description:
                continue
            unit["visual_description"] = description
            literal_text = str(unit.get("text", "")).strip()
            unit["literal_text"] = literal_text
            unit["text"] = "\n\n".join(
                part for part in (literal_text, f"Visual interpretation:\n{description}") if part
            )
            unit["has_text"] = True
            unit["extraction_method"] = f"{unit.get('extraction_method', 'ocr')}+visual"
    except ProviderError as exc:
        raise VisualUnderstandingError(str(exc)) from exc
    except (OSError, ValueError) as exc:
        raise VisualUnderstandingError("A visual page could not be prepared safely.") from exc
    return content
