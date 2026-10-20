"""OCR enrichment keeps source location and extraction provenance."""
from io import BytesIO

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend.extraction import extract_image, extract_pdf, extract_pptx


class FixedOCR:
    def __init__(self, text: str):
        self.text = text
        self.calls = 0

    def read_text(self, image) -> str:
        self.calls += 1
        return self.text


def make_png() -> bytes:
    stream = BytesIO()
    Image.new("RGB", (320, 120), "white").save(stream, format="PNG")
    return stream.getvalue()


def test_pdf_runs_ocr_when_short_native_text_may_accompany_a_visual(tmp_path):
    path = tmp_path / "workflow.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Algorithm overview")
    document.save(path)
    document.close()
    ocr = FixedOCR("Algorithm overview\nInput -> Process -> Output")

    result = extract_pdf(path, ocr)

    assert ocr.calls == 1
    assert result[0]["reference_type"] == "page"
    assert result[0]["reference_number"] == 1
    assert result[0]["native_text"] == "Algorithm overview"
    assert result[0]["ocr_text"].endswith("Input -> Process -> Output")
    assert result[0]["text"].count("Algorithm overview") == 1
    assert "Input -> Process -> Output" in result[0]["text"]
    assert result[0]["extraction_method"] == "native+ocr"
    assert result[0]["ocr_applied"] is True


def test_pdf_skips_cloud_ocr_for_clean_native_text_without_visuals(tmp_path):
    path = tmp_path / "native.pdf"
    document = pymupdf.open()
    page = document.new_page()
    native = (
        "This page contains a complete native explanation with enough selectable "
        "text to avoid a redundant paid OCR request when no visual is present."
    )
    page.insert_textbox((72, 72, 500, 300), native)
    document.save(path)
    document.close()
    ocr = FixedOCR("This should not be requested")

    result = extract_pdf(path, ocr)

    assert ocr.calls == 0
    assert result[0]["ocr_applied"] is False
    assert result[0]["extraction_method"] == "native"
    assert "complete native explanation" in result[0]["text"]


def test_pdf_scans_visual_page_even_when_native_text_is_substantial(tmp_path):
    path = tmp_path / "visual.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_textbox(
        (72, 72, 500, 300),
        "A detailed native explanation that is intentionally longer than the "
        "clean-text threshold while the page also contains a workflow diagram.",
    )
    page.draw_line((100, 350), (400, 350))
    document.save(path)
    document.close()
    ocr = FixedOCR("Start -> Process -> Finish")

    result = extract_pdf(path, ocr)

    assert ocr.calls == 1
    assert result[0]["ocr_applied"] is True
    assert "Start -> Process -> Finish" in result[0]["text"]


def test_scanned_pdf_uses_ocr_as_text_fallback(tmp_path):
    path = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    result = extract_pdf(path, FixedOCR("Text recovered from scan"))

    assert result[0]["native_text"] == ""
    assert result[0]["text"] == "Text recovered from scan"
    assert result[0]["extraction_method"] == "ocr"


def test_pptx_ocr_scans_embedded_images(tmp_path):
    image_path = tmp_path / "diagram.png"
    image_path.write_bytes(make_png())
    deck_path = tmp_path / "lesson.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Workflow"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1.5))
    deck.save(deck_path)
    ocr = FixedOCR("Start -> Validate -> Finish")

    result = extract_pptx(deck_path, ocr)

    assert ocr.calls == 1
    assert result[0]["reference_type"] == "slide"
    assert "Workflow" in result[0]["text"]
    assert "Start -> Validate -> Finish" in result[0]["text"]
    assert result[0]["ocr_applied"] is True


def test_standalone_image_has_image_reference(tmp_path):
    path = tmp_path / "notes.jpg"
    Image.new("RGB", (200, 80), "white").save(path, format="JPEG")

    result = extract_image(path, FixedOCR("Neural network diagram"))

    assert result[0]["reference_type"] == "image"
    assert result[0]["reference_number"] == 1
    assert result[0]["text"] == "Neural network diagram"
    assert result[0]["extraction_method"] == "ocr"
