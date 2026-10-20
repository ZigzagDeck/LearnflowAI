"""Smoke and regression tests for the existing upload/extract API.

All files use pytest's isolated tmp_path; no cloud credentials are required.
"""
from io import BytesIO

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

import backend.main as backend_main
from backend.processing import process_next_job
from backend.tests.fake_embeddings import TinyEmbeddingProvider


class EmptyOCR:
    def read_text(self, image):
        return ""


@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(backend_main, "UPLOAD_DIR", tmp_path)
    with TestClient(backend_main.app) as api:
        yield api


def make_pdf():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "LearnFlow lesson notes")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_pptx():
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "Machine Learning Basics"
    stream = BytesIO()
    slides.save(stream)
    return stream.getvalue()


def upload(client, name, content):
    return client.post("/materials/upload", files={"file": (name, content)})


def run_worker_once(ocr=None):
    return process_next_job(
        backend_main.get_resource_store(),
        backend_main.UPLOAD_DIR,
        ocr_factory=lambda: ocr or EmptyOCR(),
        embedding_factory=TinyEmbeddingProvider,
    )


def test_health_check(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_youtube_link_is_validated_canonicalized_and_queued(client):
    response = client.post(
        "/materials/youtube",
        json={"url": "https://youtu.be/dQw4w9WgXcQ"},
    )

    assert response.status_code == 200
    assert response.json()["source_type"] == "video"
    assert response.json()["source_url"] == (
        "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
    )
    assert response.json()["processing_job"]["status"] == "queued"

    rejected = client.post(
        "/materials/youtube",
        json={"url": "https://example.com/watch?v=dQw4w9WgXcQ"},
    )
    assert rejected.status_code == 422


def test_ocr_usage_reports_local_monthly_counter(client):
    response = client.get("/usage/ocr")

    assert response.status_code == 200
    assert response.json()["provider"] == backend_main.OCR_PROVIDER
    assert response.json()["used_pages"] == 0
    assert (
        response.json()["monthly_page_limit"]
        == backend_main.GOOGLE_DOCUMENT_AI_MONTHLY_PAGE_LIMIT
    )


def test_embedding_usage_reports_local_monthly_counter(client):
    response = client.get("/usage/embeddings")

    assert response.status_code == 200
    assert response.json()["provider"] == backend_main.EMBEDDING_PROVIDER
    assert response.json()["model"] == backend_main.VOYAGE_EMBEDDING_MODEL
    assert response.json()["dimensions"] == backend_main.VOYAGE_EMBEDDING_DIMENSIONS
    assert response.json()["used_conservative_token_units"] == 0
    assert (
        response.json()["monthly_token_unit_limit"]
        == backend_main.VOYAGE_MONTHLY_TOKEN_LIMIT
    )


def test_upload_and_extract_pdf_with_page_reference(client):
    response = upload(client, "lecture.pdf", make_pdf())
    assert response.status_code == 200
    identifier = response.json()["stored_filename"]
    resource_id = response.json()["resource_id"]
    assert identifier.endswith(".pdf")
    queued = client.post(f"/materials/{identifier}/extract")
    assert queued.status_code == 202
    run_worker_once()
    result = client.get(f"/materials/{resource_id}/content")
    assert result.status_code == 200
    assert result.json()["content"][0]["reference_type"] == "page"
    assert result.json()["content"][0]["reference_number"] == 1
    assert "LearnFlow lesson notes" in result.json()["content"][0]["text"]
    chunks = client.get(f"/materials/{resource_id}/chunks")
    assert chunks.status_code == 200
    assert chunks.json()["chunks"][0]["reference_type"] == "page"
    assert chunks.json()["chunks"][0]["reference_number"] == 1


def test_upload_and_extract_pptx_with_slide_reference(client):
    response = upload(client, "lecture.pptx", make_pptx())
    assert response.status_code == 200
    identifier = response.json()["stored_filename"]
    resource_id = response.json()["resource_id"]
    assert client.post(f"/materials/{identifier}/extract").status_code == 202
    run_worker_once()
    result = client.get(f"/materials/{resource_id}/content")
    assert result.status_code == 200
    assert result.json()["content"][0]["reference_type"] == "slide"
    assert result.json()["content"][0]["reference_number"] == 1
    assert "Machine Learning Basics" in result.json()["content"][0]["text"]


def test_duplicate_upload_reuses_existing_resource(client, tmp_path):
    content = make_pdf()
    first = upload(client, "lecture.pdf", content)
    second = upload(client, "renamed-copy.pdf", content)

    assert first.status_code == 200
    assert second.status_code == 200
    assert second.json()["deduplicated"] is True
    assert second.json()["resource_id"] == first.json()["resource_id"]
    assert second.json()["stored_filename"] == first.json()["stored_filename"]
    assert len(list(tmp_path.glob("*.pdf"))) == 1


def test_extraction_result_is_cached(client, monkeypatch):
    class CountingOCR:
        calls = 0

        def read_text(self, image):
            self.calls += 1
            return ""

    ocr = CountingOCR()
    response = upload(client, "lecture.pdf", make_pdf())
    identifier = response.json()["stored_filename"]
    resource_id = response.json()["resource_id"]

    first = client.post(f"/materials/{identifier}/extract")
    assert first.status_code == 202
    run_worker_once(ocr)
    second = client.post(f"/materials/{identifier}/extract")

    assert second.status_code == 200
    assert second.json()["cached"] is True
    assert second.json()["content"] == client.get(
        f"/materials/{resource_id}/content"
    ).json()["content"]
    assert ocr.calls == 1


def test_material_list_and_status_are_persisted(client):
    response = upload(client, "lecture.pdf", make_pdf())
    resource_id = response.json()["resource_id"]

    listing = client.get("/materials")
    status = client.get(f"/materials/{resource_id}/status")

    assert listing.status_code == 200
    assert listing.json()["materials"][0]["resource_id"] == resource_id
    assert status.status_code == 200
    assert status.json()["original_filename"] == "lecture.pdf"
    assert status.json()["status"] == "uploaded"
    assert status.json()["processing_job"]["status"] == "queued"
    assert "content_hash" not in status.json()


def test_unsupported_upload_type_is_rejected(client, tmp_path):
    response = upload(client, "malware.exe", b"not a supported type")
    assert response.status_code == 415
    assert not list(tmp_path.iterdir())


def test_upload_size_limit_and_cleanup(client, tmp_path, monkeypatch):
    monkeypatch.setattr(backend_main, "MAX_UPLOAD_SIZE_BYTES", 8)
    response = upload(client, "too-large.pdf", b"more than eight bytes")
    assert response.status_code == 413
    assert not list(tmp_path.iterdir())


def test_missing_document_returns_404(client):
    identifier = "a" * 32 + ".pdf"
    assert client.post(f"/materials/{identifier}/extract").status_code == 404


@pytest.mark.parametrize("identifier", ["invalid.pdf", "not-a-uuid", "a" * 32 + ".exe"])
def test_invalid_document_identifier(client, identifier):
    assert client.post(f"/materials/{identifier}/extract").status_code == 400


def test_malformed_pdf_is_rejected_during_extraction(client):
    response = upload(client, "broken.pdf", b"not really a PDF")
    if response.status_code == 415:
        # Hardened content validation may reject a forged PDF at upload time.
        return
    assert response.status_code == 200
    run_worker_once()
    result = client.get(f"/materials/{response.json()['resource_id']}/status")
    assert result.json()["status"] == "failed"


def test_legacy_ppt_has_explicit_conversion_error(client):
    response = upload(client, "legacy.ppt", b"old PowerPoint content")
    assert response.status_code == 415
    assert "PPTX" in response.json()["detail"]
