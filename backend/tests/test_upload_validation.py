"""Upload content-check regressions; no network or cloud dependencies."""
from io import BytesIO
import pytest
import pymupdf
from pptx import Presentation
from PIL import Image
from fastapi.testclient import TestClient
import backend.main as api
from backend.processing import process_next_job


class EmptyOCR:
    def read_text(self, image):
        return ""


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "UPLOAD_DIR", tmp_path)
    return TestClient(api.app), tmp_path


def post(client, filename, data):
    return client.post("/materials/upload", files={"file": (filename, data)})


def run_worker_once(directory, ocr=None):
    return process_next_job(
        api.get_resource_store(),
        directory,
        ocr_factory=lambda: ocr or EmptyOCR(),
    )


def pdf_bytes():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Verified lesson PDF")
    raw = doc.tobytes()
    doc.close()
    return raw


def pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide with citation"
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


def image_bytes(format_name="PNG"):
    stream = BytesIO()
    Image.new("RGB", (120, 60), "white").save(stream, format=format_name)
    return stream.getvalue()


@pytest.mark.parametrize("filename,data", [
    ("fake.pdf", b"not a pdf"),
    ("fake.pptx", b"not a zip"),
    ("fake.ppt", b"not OLE"),
    ("fake.avi", b"RIFF0000WAVE"),
    ("fake.mp4", b"broken ftyp"),
    ("fake.mov", b"invalid"),
    ("fake.mkv", b"invalid"),
    ("fake.webm", b"invalid"),
])
def test_rejects_fake_upload_and_deletes_it(client, filename, data):
    api_client, directory = client
    response = post(api_client, filename, data)
    assert response.status_code == 415, response.text
    assert list(directory.iterdir()) == []


def test_rejects_empty_upload(client):
    api_client, directory = client
    response = post(api_client, "empty.pdf", b"")
    assert response.status_code == 422
    assert list(directory.iterdir()) == []


@pytest.mark.parametrize("ext,contents,kind", [
    ("pdf", pdf_bytes, "page"),
    ("pptx", pptx_bytes, "slide"),
])
def test_upload_and_extract_real_content(client, ext, contents, kind):
    api_client, directory = client
    response = post(api_client, f"lesson.{ext}", contents())
    assert response.status_code == 200, response.text
    source = response.json()["stored_filename"]
    resource_id = response.json()["resource_id"]
    assert (directory / source).is_file()
    assert api_client.post(f"/materials/{source}/extract").status_code == 202
    run_worker_once(directory)
    output = api_client.get(f"/materials/{resource_id}/content")
    assert output.status_code == 200, output.text
    assert output.json()["content"][0]["reference_type"] == kind
    assert output.json()["content"][0]["reference_number"] == 1


@pytest.mark.parametrize("filename,format_name", [
    ("diagram.png", "PNG"),
    ("diagram.jpg", "JPEG"),
    ("diagram.jpeg", "JPEG"),
])
def test_accepts_supported_images(client, filename, format_name, monkeypatch):
    api_client, directory = client

    class ImageOCR:
        def read_text(self, image):
            return "Input Validate Output"

    response = post(api_client, filename, image_bytes(format_name))
    assert response.status_code == 200, response.text
    source = response.json()["stored_filename"]
    resource_id = response.json()["resource_id"]
    assert (directory / source).is_file()
    run_worker_once(directory, ImageOCR())
    output = api_client.get(f"/materials/{resource_id}/content")
    assert output.status_code == 200, output.text
    assert output.json()["content"][0]["reference_type"] == "image"
    assert output.json()["content"][0]["text"] == "Input Validate Output"


def test_rejects_image_with_mismatched_extension(client):
    api_client, directory = client
    response = post(api_client, "diagram.png", image_bytes("JPEG"))
    assert response.status_code == 415
    assert list(directory.iterdir()) == []
