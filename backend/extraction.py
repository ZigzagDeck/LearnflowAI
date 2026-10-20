from pathlib import Path
from hashlib import sha256
from typing import Callable

import pymupdf
from pptx import Presentation

from .image_utils import load_rgb_image
from .ocr import OCRReader


MAX_EXTRACTION_UNITS = 500
MAX_EXTRACTED_CHARACTERS = 2_000_000
PDF_OCR_DPI = 180
MAX_RENDERED_PAGE_PIXELS = 20_000_000
MIN_CLEAN_NATIVE_TEXT_CHARACTERS = 80
ProgressCallback = Callable[[int, int], None]


def _merge_native_and_ocr(native_text: str, ocr_text: str) -> str:
    native_text = native_text.strip()
    native_normalized = " ".join(native_text.casefold().split())
    additions: list[str] = []
    seen: set[str] = set()
    for line in ocr_text.splitlines():
        clean = " ".join(line.split())
        normalized = clean.casefold()
        if not normalized or normalized in seen or normalized in native_normalized:
            continue
        seen.add(normalized)
        additions.append(clean)
    return "\n".join(part for part in (native_text, "\n".join(additions)) if part)


def _render_pdf_page(page: pymupdf.Page) -> bytes:
    zoom = PDF_OCR_DPI / 72
    estimated_pixels = page.rect.width * zoom * page.rect.height * zoom
    if estimated_pixels > MAX_RENDERED_PAGE_PIXELS:
        zoom *= (MAX_RENDERED_PAGE_PIXELS / estimated_pixels) ** 0.5
    pixmap = page.get_pixmap(
        matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csRGB, alpha=False
    )
    return pixmap.tobytes("png")


def _pdf_page_requires_ocr(page: pymupdf.Page, native_text: str) -> bool:
    """OCR scans, short pages, and pages containing raster/vector visuals."""
    if len(" ".join(native_text.split())) < MIN_CLEAN_NATIVE_TEXT_CHARACTERS:
        return True
    return bool(page.get_images(full=True) or page.get_drawings())


def _content_record(
    reference_type: str,
    reference_number: int,
    native_text: str,
    ocr_text: str,
    *,
    ocr_applied: bool,
) -> dict[str, int | str | bool]:
    text = _merge_native_and_ocr(native_text, ocr_text)
    if native_text and ocr_text:
        method = "native+ocr"
    elif native_text:
        method = "native"
    else:
        method = "ocr"
    return {
        "reference_type": reference_type,
        "reference_number": reference_number,
        "text": text,
        "native_text": native_text.strip(),
        "ocr_text": ocr_text.strip(),
        "has_text": bool(text),
        "ocr_applied": ocr_applied,
        "extraction_method": method,
    }


def extract_pdf(
    path: Path, ocr: OCRReader, progress: ProgressCallback | None = None
) -> list[dict[str, int | str | bool]]:
    pages: list[dict[str, int | str | bool]] = []
    character_count = 0

    with pymupdf.open(path) as document:
        if document.page_count > MAX_EXTRACTION_UNITS:
            raise ValueError(
                f"This PDF has too many pages. The maximum is {MAX_EXTRACTION_UNITS} pages."
            )

        for page_number, page in enumerate(document, start=1):
            native_text = page.get_text("text").strip()
            should_scan = _pdf_page_requires_ocr(page, native_text)
            ocr_text = ocr.read_text(_render_pdf_page(page)) if should_scan else ""
            record = _content_record(
                "page",
                page_number,
                native_text,
                ocr_text,
                ocr_applied=should_scan,
            )
            text = str(record["text"])
            character_count += len(text)
            if character_count > MAX_EXTRACTED_CHARACTERS:
                raise ValueError(
                    "This document contains too much text to extract in one request."
                )
            pages.append(record)
            if progress:
                progress(page_number, document.page_count)

    return pages


def extract_pptx(
    path: Path, ocr: OCRReader, progress: ProgressCallback | None = None
) -> list[dict[str, int | str | bool]]:
    slides: list[dict[str, int | str | bool]] = []
    character_count = 0
    presentation = Presentation(path)

    if len(presentation.slides) > MAX_EXTRACTION_UNITS:
        raise ValueError(
            f"This presentation has too many slides. The maximum is {MAX_EXTRACTION_UNITS} slides."
        )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        native_text = "\n".join(text_parts)
        ocr_parts: list[str] = []
        seen_images: set[str] = set()
        for shape in slide.shapes:
            image = getattr(shape, "image", None)
            if image is None:
                continue
            blob = image.blob
            digest = sha256(blob).hexdigest()
            if digest in seen_images:
                continue
            seen_images.add(digest)
            # Decode before inference to reject oversized or corrupt embedded images.
            ocr_parts.append(ocr.read_text(load_rgb_image(blob)))
        ocr_text = "\n".join(part for part in ocr_parts if part)
        record = _content_record(
            "slide", slide_number, native_text, ocr_text,
            ocr_applied=bool(seen_images),
        )
        text = str(record["text"])
        character_count += len(text)
        if character_count > MAX_EXTRACTED_CHARACTERS:
            raise ValueError(
                "This presentation contains too much text to extract in one request."
            )
        slides.append(record)
        if progress:
            progress(slide_number, len(presentation.slides))

    return slides


def extract_image(
    path: Path, ocr: OCRReader, progress: ProgressCallback | None = None
) -> list[dict[str, int | str | bool]]:
    text = ocr.read_text(path)
    if progress:
        progress(1, 1)
    return [_content_record("image", 1, "", text, ocr_applied=True)]
